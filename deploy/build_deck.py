"""
Build the Nu Score On the Go product deck (16:9, Northern Arc / AltiFi palette).

Usage: .venv/bin/python deploy/build_deck.py [result.json] [out.pptx]
  result.json  a training result (defaults to the latest job under DATA_ROOT)
  out.pptx     defaults to docs/Nu_Score_On_the_Go.pptx
Screenshots are read from docs/media/ when present.
"""

from __future__ import annotations

import json
import os
import sys
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
MEDIA = ROOT / "docs" / "media"
BRAND = ROOT / "frontend" / "brand"

NAVY = RGBColor(0x0D, 0x1C, 0x31)
NAVY2 = RGBColor(0x16, 0x29, 0x4A)
CYAN = RGBColor(0x4F, 0xC6, 0xE0)
CYAN_DEEP = RGBColor(0x1B, 0x9A, 0xB8)
WASH = RGBColor(0xE9, 0xF8, 0xFC)
INK = RGBColor(0x0D, 0x1C, 0x31)
INK2 = RGBColor(0x44, 0x54, 0x6B)
INK3 = RGBColor(0x7A, 0x87, 0x9A)
LINE = RGBColor(0xE3, 0xE8, 0xEF)
GROUND = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOOD = RGBColor(0x1F, 0x9D, 0x6B)
WARN = RGBColor(0xE9, 0xB9, 0x49)
BAD = RGBColor(0xC8, 0x45, 0x3B)
MUTED = RGBColor(0x8F, 0xA3, 0xBF)
BAND_COLORS = {"A+": GOOD, "A": CYAN, "B": MUTED, "C": WARN, "D": RGBColor(0xE0, 0x7B, 0x39), "E": BAD}
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)          # side margin


# ------------------------------------------------------------------ helpers


def rgb_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def text(slide, left, top, width, height, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.1):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts}) runs."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paragraphs = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        pieces = para if isinstance(para, list) else [(para, {})]
        for piece, opts in pieces:
            r = p.add_run()
            r.text = piece
            f = r.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.color.rgb = opts.get("color", color)
    return box


def rect(slide, left, top, width, height, color, shape=MSO_SHAPE.RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    rgb_fill(s, color)
    if line is not None:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def card(slide, left, top, width, height, title, body, accent=CYAN, title_size=14, body_size=11.5, fill=WHITE):
    s = rect(slide, left, top, width, height, fill, line=LINE)
    rect(slide, left, top + Inches(0.12), Inches(0.05), height - Inches(0.24), accent)
    text(slide, left + Inches(0.18), top + Inches(0.08), width - Inches(0.3), Inches(0.4), title, size=title_size, bold=True)
    text(slide, left + Inches(0.18), top + Inches(0.48), width - Inches(0.3), height - Inches(0.55), body, size=body_size, color=INK2)
    return s


def bullets(slide, left, top, width, height, items, size=12.5, color=INK2, gap=1.15):
    paras = []
    for item in items:
        if isinstance(item, tuple):
            paras.append([(item[0], {"bold": True, "color": INK}), (item[1], {})])
        else:
            paras.append([("▪  ", {"color": CYAN_DEEP, "bold": True}), (item, {})])
    return text(slide, left, top, width, height, paras, size=size, color=color, line_spacing=gap)


def _duration(seconds: float) -> str:
    seconds = float(seconds or 0)
    if seconds < 90:
        return f"{seconds:.0f} seconds"
    return f"about {seconds / 60:.0f} minutes"


def _chevrons(slide, top, steps, color=CYAN):
    sw = (W - 2 * M) / len(steps)
    for i, (title, body) in enumerate(steps):
        left = M + i * sw
        rect(slide, left, top, sw - Inches(0.08), Inches(0.5), color if i % 2 == 0 else CYAN_DEEP, shape=MSO_SHAPE.CHEVRON)
        text(slide, left + Inches(0.25), top + Inches(0.02), sw - Inches(0.5), Inches(0.45), title, size=11, bold=True,
             color=NAVY if i % 2 == 0 else WHITE, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, left + Inches(0.05), top + Inches(0.58), sw - Inches(0.2), Inches(0.6), body, size=10, color=INK2)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def slide(self, title=None, subtitle=None, dark=False, section=None):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bg = rect(s, 0, 0, W, H, NAVY if dark else WHITE)
        if not dark:
            rect(s, 0, 0, W, Inches(0.08), NAVY)
            rect(s, 0, Inches(0.08), Inches(2.2), Inches(0.04), CYAN)
        if title:
            text(s, M, Inches(0.3), W - 2 * M, Inches(1.0), title, size=21, bold=True, color=WHITE if dark else INK)
        if subtitle:
            text(s, M, Inches(1.32), W - 2 * M, Inches(0.4), subtitle, size=12, color=CYAN if dark else INK3)
        # footer
        foot = WHITE if dark else INK3
        text(s, M, H - Inches(0.42), Inches(8), Inches(0.3), "Nu Score On the Go  ·  Northern Arc / AltiFi  ·  Internal",
             size=9, color=foot)
        text(s, W - M - Inches(1.0), H - Inches(0.42), Inches(1.0), Inches(0.3), str(self.n), size=9, color=foot, align=PP_ALIGN.RIGHT)
        if section:
            text(s, W - M - Inches(4), Inches(0.14), Inches(4), Inches(0.3), section.upper(), size=8.5, color=CYAN_DEEP, align=PP_ALIGN.RIGHT, bold=True)
        return s

    def logos(self, s, top, light=True):
        na = BRAND / ("northern-arc-light.png" if light else "northern-arc.png")
        al = BRAND / ("altifi_logo_light.png" if light else "altifi_logo.png")
        if na.exists():
            s.shapes.add_picture(str(na), M, top, height=Inches(0.55))
        if al.exists():
            s.shapes.add_picture(str(al), M + Inches(1.35), top + Inches(0.02), height=Inches(0.5))

    def picture(self, s, path, left, top, width=None, height=None, border=True):
        """Place an image inside a (width × height) box, preserving aspect ratio;
        the image is centred horizontally in the box when it is narrower."""
        if not Path(path).exists():
            rect(s, left, top, width or Inches(5), height or Inches(3), GROUND, line=LINE)
            text(s, left, top + Inches(1), width or Inches(5), Inches(0.5), f"[screenshot: {Path(path).name}]",
                 size=10, color=INK3, align=PP_ALIGN.CENTER)
            return None
        if width and height:
            from PIL import Image
            with Image.open(path) as im:
                ratio = im.size[0] / im.size[1]
            if width / height > ratio:            # box is wider than the image: fit height
                new_w = int(height * ratio)
                left = left + int((width - new_w) / 2)
                width, height = new_w, height
            else:
                height = int(width / ratio)
        pic = s.shapes.add_picture(str(path), left, top, width=width, height=height)
        if border:
            pic.line.color.rgb = LINE
            pic.line.width = Pt(0.75)
        return pic

    def save(self, path):
        self.prs.save(str(path))


# --------------------------------------------------------------------- deck


def build(result: dict, out: Path) -> None:
    d = Deck()
    ds = result.get("dataset", {})
    lb = result.get("leaderboard", [])
    nu = result.get("nuscore", {})
    champ = result.get("champion", {})
    n_models = len(lb)
    today = date.today().strftime("%d %B %Y")

    # 1 · Title ---------------------------------------------------------------
    s = d.slide(dark=True)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, Inches(5.9), W, Inches(1.6), NAVY2)
    for i in range(14):
        rect(s, Inches(0.2 + i * 0.95), Inches(0.2), Emu(12000), Inches(5.6), RGBColor(0x14, 0x2A, 0x44))
    d.logos(s, Inches(0.55))
    text(s, M, Inches(2.0), Inches(9), Inches(0.5), "NU SCORE ON THE GO", size=13, bold=True, color=CYAN)
    text(s, M, Inches(2.45), Inches(11), Inches(1.6), "Build ML models with your data", size=44, bold=True, color=WHITE)
    text(s, M, Inches(3.75), Inches(10), Inches(1.0),
         "From loan book, bureau and bank-statement files to a validated, explainable Nu Score — "
         "and live scoring of new applications — in minutes, on Northern Arc's own servers.",
         size=16, color=RGBColor(0xC7, 0xD2, 0xE0))
    text(s, M, Inches(6.2), Inches(8), Inches(0.4), f"Product overview  ·  {today}  ·  Internal", size=12, color=CYAN)
    text(s, W - M - Inches(4.5), Inches(6.2), Inches(4.5), Inches(0.9),
         "http://10.91.16.86/nuscore_claudecode/", size=12, color=WHITE, align=PP_ALIGN.RIGHT)

    # 2 · Executive summary ----------------------------------------------------
    s = d.slide("Nu Score On the Go turns the data a lender already has into a governed credit score in minutes",
                "Executive summary", section="Executive summary")
    cols = [
        ("The problem", CYAN_DEEP, [
            "Building a credit model takes a data-science team weeks: joining sources, cleaning formats, "
            "choosing algorithms, validating honestly, packaging for underwriters.",
            "Most lending teams never get beyond a spreadsheet — decisions stay judgemental and inconsistent.",
        ]),
        ("What we built", CYAN, [
            "Upload the files you have; the engine matches identifiers, screens for PII and leakage, "
            "tunes nine model families on GPUs and selects a champion under model-risk discipline.",
            "Output is a calibrated 0–1000 Nu Score with A+…E bands, reason codes and reports — "
            "and the model scores new applications immediately.",
        ]),
        ("Why it matters", GOOD, [
            f"Live today on internal infrastructure; a {ds.get('n_rows', 0):,}-row book trains "
            f"{n_models} models in {_duration(ds.get('training_seconds', 0))}.",
            "Data never leaves the host and is deleted once results are delivered; no PII required.",
            "Free during introduction, then ₹10 per case scored — compute included.",
        ]),
    ]
    cw = (W - 2 * M - Inches(0.4)) / 3
    for i, (title, accent, items) in enumerate(cols):
        left = M + i * (cw + Inches(0.2))
        card(s, left, Inches(1.75), cw, Inches(3.9), title, "", accent=accent, title_size=15)
        bullets(s, left + Inches(0.2), Inches(2.25), cw - Inches(0.35), Inches(3.3), items, size=12)
    # KPI strip
    kpis = [("9", "model families incl. FT-Transformer"), ("5", "validation cohorts"), ("0–1000", "Nu Score, 6 bands"),
            ("20 GB", "per file, resumable"), ("₹10", "per case after launch")]
    kw = (W - 2 * M) / len(kpis)
    for i, (big, small) in enumerate(kpis):
        left = M + i * kw
        rect(s, left, Inches(5.85), kw - Inches(0.1), Inches(0.95), GROUND)
        text(s, left, Inches(5.9), kw - Inches(0.1), Inches(0.5), big, size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        text(s, left, Inches(6.35), kw - Inches(0.1), Inches(0.4), small, size=10, color=INK3, align=PP_ALIGN.CENTER)

    # 3 · The problem ----------------------------------------------------------
    s = d.slide("Today a usable credit model is a six-week project; the product collapses it to one sitting",
                "The problem we solve", section="Context")
    steps_old = [("Week 1–2", "Collect & join loan, bureau, bank files by hand"), ("Week 2–3", "Clean formats, chase missing identifiers"),
                 ("Week 3–4", "Try 2–3 algorithms in notebooks"), ("Week 4–5", "Validate, argue about leakage & splits"),
                 ("Week 5–6", "Package for underwriting; write the committee pack")]
    steps_new = [("Minute 1", "Drop files; identifiers matched"), ("Minute 2", "PII / leakage screened; outcome confirmed"),
                 ("Minute 3–6", "9 families tuned on GPU; champion picked"), ("Minute 6", "Nu Score bands, SHAP, warnings, report"),
                 ("Minute 7", "Score the next application")]
    for row, (label, steps, color) in enumerate([("Conventional build", steps_old, MUTED), ("Nu Score On the Go", steps_new, CYAN)]):
        top = Inches(2.0 + row * 2.5)
        text(s, M, top, Inches(3), Inches(0.4), label, size=14, bold=True, color=NAVY)
        sw = (W - 2 * M) / len(steps)
        for i, (when, what) in enumerate(steps):
            left = M + i * sw
            chev = rect(s, left, top + Inches(0.5), sw - Inches(0.08), Inches(0.55), color, shape=MSO_SHAPE.CHEVRON)
            text(s, left + Inches(0.25), top + Inches(0.55), sw - Inches(0.5), Inches(0.45), when, size=11, bold=True,
                 color=WHITE if color is not CYAN else NAVY, anchor=MSO_ANCHOR.MIDDLE)
            text(s, left + Inches(0.05), top + Inches(1.15), sw - Inches(0.2), Inches(0.9), what, size=11, color=INK2)
    text(s, M, Inches(6.55), W - 2 * M, Inches(0.4),
         "Quick tier on a 20,000-row book: about 3 minutes end to end on the current host (2 × Tesla T4 for training).",
         size=10.5, color=INK3)

    # 4 · Product at a glance --------------------------------------------------
    s = d.slide("One guided flow from raw files to a scored application", "Product at a glance", section="Product")
    flow = [("01", "Upload sources", "CSV, bureau JSON, Parquet; up to 20 GB; resumable"),
            ("02", "Match identifiers", "Roles proposed by a local LLM; keys verified by value overlap"),
            ("03", "Analyse & screen", "PII, leakage, IDs flagged; IV and WOE bins; clean-up log"),
            ("04", "Train model zoo", "Nine families tuned with Optuna on GPUs"),
            ("05", "Pick champion", "Composite on selection-safe cohorts; stacked blend"),
            ("06", "Nu Score bands", "Calibrated 0–1000 score; bands learned from data"),
            ("07", "Score applications", "Form or batch file; reason codes; drift check")]
    fw = (W - 2 * M) / len(flow)
    for i, (num, title, body) in enumerate(flow):
        left = M + i * fw
        rect(s, left, Inches(1.9), fw - Inches(0.1), Inches(0.08), CYAN if i % 2 == 0 else CYAN_DEEP)
        text(s, left, Inches(2.05), fw, Inches(0.4), num, size=11, bold=True, color=CYAN_DEEP)
        text(s, left, Inches(2.4), fw - Inches(0.15), Inches(0.7), title, size=14, bold=True, color=NAVY)
        text(s, left, Inches(3.05), fw - Inches(0.2), Inches(1.6), body, size=11, color=INK2)
    d.picture(s, MEDIA / "landing_hero.png", M, Inches(4.1), width=W - 2 * M, height=Inches(2.75))

    # 5 · Inputs ---------------------------------------------------------------
    s = d.slide("Bring what you have: a loan dump plus any bureau and bank-statement variables, joined by one identifier",
                "Inputs", section="Product")
    cw = (W - 2 * M - Inches(0.4)) / 3
    cards = [
        ("REQUIRED · Loan dump", GOOD, ["loan_id / account_no", "Outcome to predict: default flag, DPD bucket, write-off, closure",
                                        "Product, amount, tenure, rate, EMI, LTV, disbursal date", "Repayment history: DPD, overdue, months on book",
                                        "Applicant profile without PII: age band, income, occupation, pincode"]),
        ("RECOMMENDED · Bureau", CYAN, ["Bureau score, enquiries (3/6/12 m)", "Live & closed trade lines, exposure, utilisation",
                                        "Worst DPD, write-offs, settlements, vintage", "Flat CSV or the raw bureau JSON (nested trade lines flattened)"]),
        ("RECOMMENDED · Bank statement", CYAN, ["Average / minimum / month-end balances", "Credit & debit counts and amounts",
                                                "Salary regularity, cheque or EMI bounces", "Cash withdrawals, existing EMI outflow; monthly rows are fine"]),
    ]
    for i, (title, accent, items) in enumerate(cards):
        left = M + i * (cw + Inches(0.2))
        card(s, left, Inches(1.75), cw, Inches(3.35), title, "", accent=accent, title_size=13)
        bullets(s, left + Inches(0.2), Inches(2.2), cw - Inches(0.35), Inches(2.8), items, size=11)
    rules = [("Identifier  ", "one common key across files; names may differ (loan_id, ref_no, account_number) — the engine finds the match and shows the match rate"),
             ("Outcome  ", "0/1, yes/no, a class label, or a number to threshold (e.g. DPD > 90); you confirm the positive class"),
             ("Time column  ", "disbursal or snapshot date unlocks true out-of-time validation"),
             ("Keep out  ", "names, mobile, email, address, PAN, Aadhaar, DOB — flagged and dropped if present; not needed for the score")]
    bullets(s, M, Inches(5.25), W - 2 * M, Inches(1.5), rules, size=11)

    # 6 · Intake & screening ---------------------------------------------------
    s = d.slide("Identifiers are matched by measured overlap and every column is screened before a model sees it",
                "Data intake and screening", section="How it works")
    left_items = [("Matching  ", "a local language model proposes file roles (loan dump / bureau / bank statement); join keys are chosen by "
                                  "actual value overlap between candidate columns and shown with match rate and rows per case"),
                  ("Bureau JSON  ", "nested objects become columns; trade-line lists become count / sum / max / min per field"),
                  ("Screening  ", "PII by name and value pattern (emails, mobiles, PAN, Aadhaar); leakage by name and implausible IV; "
                                  "identifiers by uniqueness — all excluded by default, overridable"),
                  ("Clean-up  ", "₹ / $ symbols, % signs, 1,23,456 lakh and 1,234,567 grouping, yes/no, NA markers — normalised deterministically "
                                 "and logged for review"),
                  ("Profiling  ", "DuckDB summarises millions of rows in seconds: types, nulls, distinct counts, date detection, IV with WOE bins")]
    bullets(s, M, Inches(1.8), Inches(6.6), Inches(4.8), left_items, size=11.5)
    tbl_left = M + Inches(6.9)
    rows = [("Check", "Trigger", "Action"), ("PII", "name pattern or value pattern", "dropped before training"),
            ("Leakage", "post-outcome name, IV > 1.0", "excluded by default"), ("Identifier", "> 95 % unique", "kept as key, not a feature"),
            ("Constant / missing", "one value or all null", "listed in diagnostics"), ("Correlated pairs", "|r| ≥ 0.8", "listed; importances shared")]
    tb = s.shapes.add_table(len(rows), 3, tbl_left, Inches(1.85), W - M - tbl_left, Inches(2.6)).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tb.cell(r, c)
            cell.text = val
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(10.5); para.font.name = FONT
            para.font.bold = r == 0
            para.font.color.rgb = WHITE if r == 0 else INK2
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if r == 0 else (GROUND if r % 2 else WHITE)
    text(s, tbl_left, Inches(4.7), W - M - tbl_left, Inches(1.8),
         [[("Example from a live run: ", {"bold": True, "color": INK})],
          [("loan_id = ref_no  ·  92 % of cases matched  ·  1 row per case (bureau)", {"color": INK2})],
          [("loan_id = account_number  ·  85 % matched  ·  6 rows per case (bank statement, latest kept)", {"color": INK2})],
          [("customer_name → PII · mobile → PII · writeoff_amount → leakage (IV 13.9)", {"color": INK2})]], size=11)

    # 7 · Engine --------------------------------------------------------------
    s = d.slide("Nine model families are tuned automatically; a stacked blend adds the AI-enhanced probability",
                "The modelling engine", section="How it works")
    zoo = [("Logistic regression", "Interpretable baseline every scorecard starts from"),
           ("Random Forest / Extra Trees", "Bagged trees; robust to outliers and monotone transforms"),
           ("Histogram Gradient Boosting", "scikit-learn boosting with native missing handling"),
           ("XGBoost (GPU) / LightGBM", "Regularised gradient boosting; fast on wide data"),
           ("CatBoost", "Ordered boosting; strong on small data"),
           ("MLP", "Neural network on standardised features"),
           ("FT-Transformer (PyTorch)", "Attention across tabular features; trained on the T4s"),
           ("NuScore Blend", "Logistic meta-learner on out-of-fold predictions of the top four")]
    for i, (name, desc) in enumerate(zoo):
        top = Inches(1.8 + i * 0.58)
        rect(s, M, top, Inches(0.08), Inches(0.45), CYAN if i < 7 else NAVY)
        text(s, M + Inches(0.2), top - Inches(0.02), Inches(3.2), Inches(0.5), name, size=12, bold=True, color=NAVY)
        text(s, M + Inches(3.4), top - Inches(0.02), Inches(3.6), Inches(0.5), desc, size=10.5, color=INK2)
    right = M + Inches(7.4)
    card(s, right, Inches(1.8), W - M - right, Inches(2.3), "Tuning", "", accent=CYAN)
    bullets(s, right + Inches(0.2), Inches(2.25), W - M - right - Inches(0.35), Inches(1.9),
            ["Optuna TPE search with median pruning and early stopping per family",
             "Quick 4 trials / 2 folds  ·  Balanced 25 / 3  ·  Thorough 60 / 3",
             "Tuning on ≤ 200 k rows; final fit on ≤ 1 M rows; evaluation on all rows loaded"], size=10.5)
    card(s, right, Inches(4.3), W - M - right, Inches(2.3), "Preprocessing", "", accent=CYAN_DEEP)
    bullets(s, right + Inches(0.2), Inches(4.75), W - M - right - Inches(0.35), Inches(1.9),
            ["Median imputation with missing-value indicators",
             "One-hot for low-cardinality categoricals; out-of-fold target encoding for pincode-like columns",
             "Dates → month and days-to-reference; free text dropped; schema saved for scoring"], size=10.5)

    # 8 · Validation discipline -------------------------------------------------
    s = d.slide("Five cohorts keep selection honest: test and out-of-time never influence which model wins",
                "Validation discipline", section="How it works")
    cohorts = [("Train", 55, "fitting & tuning", CYAN_DEEP), ("Validation", 12, "selection, thresholds", CYAN),
               ("Calibration", 10, "Platt scaling, band cut-offs", MUTED), ("Test", 10, "untouched estimate + bootstrap CI", NAVY2),
               ("Out-of-time", 13, "latest dates; rows sharing a date never split", NAVY)]
    total_w = W - 2 * M
    left = M
    for name, share, job, color in cohorts:
        wdt = int(total_w * share / 100)
        rect(s, left, Inches(1.85), wdt - Inches(0.04), Inches(1.15), color)
        fg = NAVY if color in (CYAN, MUTED) else WHITE
        text(s, left, Inches(1.9), wdt - Inches(0.04), Inches(0.3), name, size=11, bold=True, color=fg, align=PP_ALIGN.CENTER)
        text(s, left, Inches(2.17), wdt - Inches(0.04), Inches(0.3), f"{share}%", size=11, color=fg, align=PP_ALIGN.CENTER)
        text(s, left, Inches(2.45), wdt - Inches(0.04), Inches(0.55), job, size=8, color=fg, align=PP_ALIGN.CENTER)
        left += wdt
    rect(s, M, Inches(3.1), int(total_w * 0.77), Inches(0.06), CYAN)
    text(s, M, Inches(3.17), int(total_w * 0.77), Inches(0.35), "used for selection", size=10, color=CYAN_DEEP, align=PP_ALIGN.CENTER)
    rect(s, M + int(total_w * 0.77), Inches(3.1), int(total_w * 0.23), Inches(0.06), BAD)
    text(s, M + int(total_w * 0.77), Inches(3.17), int(total_w * 0.23), Inches(0.35), "never used for selection", size=10, color=BAD, align=PP_ALIGN.CENTER)
    items = [("Champion composite  ", "validation AUC 45 % · validation Brier 20 % · validation→calibration consistency 15 % · overfit gap 10 % · simplicity 10 %"),
             ("Chronological  ", "when a time column exists, cohorts follow observation dates and the OOT cohort is the latest period; "
                                 "otherwise it is a random hold-out and the UI says so"),
             ("Calibration  ", "Platt scaling fitted on the calibration cohort — smooth, monotone, never saturating"),
             ("Warnings engine  ", "weak AUC, overfit gap > 0.10, validation→OOT drop, calibration worsening Brier, duplicate rows, correlated pairs, small cohorts"),
             ("Unit tests  ", "prove that flipping test / OOT labels cannot change the champion")]
    bullets(s, M, Inches(3.75), W - 2 * M, Inches(3.0), items, size=12.5, gap=1.3)

    # 9 · Champion / results ---------------------------------------------------
    s = d.slide(f"Live run: {champ.get('name', 'the champion')} won on selection cohorts and held up on the untouched test and OOT cohorts",
                f"Results from a training run  ·  {ds.get('n_rows', 0):,} rows  ·  {n_models} models  ·  tier: {ds.get('tier', '-')}",
                section="Results")
    d.picture(s, MEDIA / "champion.png", M, Inches(1.7), width=Inches(7.2), height=Inches(2.7))
    if lb:
        cd = CategoryChartData()
        cd.categories = [m["model"] for m in lb]
        cd.add_series("Validation AUC", [round(m.get("validation_auc", 0), 3) for m in lb])
        cd.add_series("Test AUC", [round(m.get("roc_auc", 0), 3) for m in lb])
        cd.add_series("OOT AUC", [round(m.get("oot_auc", 0), 3) for m in lb])
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, M + Inches(7.4), Inches(1.6), W - M - M - Inches(7.4), Inches(5.2), cd)
        ch = gf.chart
        ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(9)
        ch.category_axis.tick_labels.font.size = Pt(9)
        ch.category_axis.reverse_order = True
        ch.value_axis.minimum_scale = 0.5; ch.value_axis.maximum_scale = 1.0
        ch.value_axis.tick_labels.font.size = Pt(9)
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.major_gridlines.format.line.color.rgb = LINE
        for series, color in zip(ch.series, (CYAN, NAVY, MUTED)):
            series.format.fill.solid(); series.format.fill.fore_color.rgb = color
        ch.plots[0].gap_width = 60
    text(s, M, Inches(4.55), Inches(7.2), Inches(2.2),
         [[("Composite breakdown ", {"bold": True, "color": INK}), ("shows why the champion won; the leaderboard ranks every model on the same "
           "composite and reports validation, test and OOT AUC side by side so a reviewer can see stability at a glance.", {})],
          [("Bootstrap 95 % interval ", {"bold": True, "color": INK}), ("on the test AUC quantifies sampling uncertainty; the OOT tile states whether it is a true out-of-time cohort.", {})]],
         size=11, color=INK2)

    # 10 · Nu Score ----------------------------------------------------------------
    s = d.slide("The Nu Score is a calibrated 0–1000 scale with bands learned from the data, not fixed cut-offs",
                "Nu Score (AI-enhanced)", section="Results")
    bullets(s, M, Inches(1.75), Inches(5.6), Inches(2.6),
            [("Scale  ", "0–1000, higher = safer; 500 = even odds; +60 points halves the odds of going bad"),
             ("Calibrated  ", "Platt scaling on the calibration cohort, so the score tracks observed bad rates"),
             ("AI-enhanced  ", "probability from the stacked blend; base score and AI adjustment shown separately for governance"),
             ("Bands  ", "A+ … E cut-offs learned on the calibration cohort: quantile seeds merged until bad rate rises monotonically and each band holds ≥ 5 %"),
             ("Cut-off simulator  ", "approval rate vs bad rate among approved vs bads rejected, on the test cohort")], size=11)
    bt = (nu.get("band_table") or {}).get("test") or []
    if bt:
        rows = [("Band", "Score range", "Share", "Bad rate", "Lift")] + [
            (b["band"], f"{b['min_score']}–{b['max_score']}", f"{100 * b['share']:.1f}%", f"{100 * (b['bad_rate'] or 0):.1f}%",
             f"{b['lift']:.2f}×" if b.get("lift") is not None else "–") for b in bt]
        tb = s.shapes.add_table(len(rows), 5, M, Inches(4.45), Inches(5.6), Inches(2.3)).table
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = tb.cell(r, c); cell.text = val
                para = cell.text_frame.paragraphs[0]; para.font.size = Pt(10); para.font.name = FONT
                para.font.bold = r == 0 or c == 0
                para.font.color.rgb = WHITE if r == 0 else INK2
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY if r == 0 else (BAND_COLORS.get(row[0], GROUND) if c == 0 else WHITE)
                if r > 0 and c == 0:
                    para.font.color.rgb = WHITE if row[0] in ("A+", "D", "E") else NAVY
    d.picture(s, MEDIA / "nuscore.png", M + Inches(5.9), Inches(1.7), width=W - 2 * M - Inches(5.9), height=Inches(5.1))

    # 11 · Explainability & governance -----------------------------------------
    s = d.slide("Every score is explainable and every run carries its own validation notes, diagnostics and drift check",
                "Explainability and governance", section="Results")
    cw = (W - 2 * M - Inches(0.4)) / 3
    gov = [("Explain", CYAN, ["SHAP drivers, beeswarm and dependence plots for the champion", "Per-record reason codes (adverse-action style) with a plain-English glossary",
                              "Permutation importance where a model has no native importances", "AI-written narrative from aggregate metrics only"]),
           ("Validate", CYAN_DEEP, ["Amber validation notes generated from the data", "Bootstrap AUC interval, duplicate-row and correlation checks",
                                    "WOE bin tables per variable; IV strength flags", "Calibration reliability before vs after"]),
           ("Monitor", GOOD, ["PSI drift of score and every feature per cohort vs training", "Drift check on every batch scored later",
                              "Rows without an outcome scored separately, never in a cohort", "HTML / PDF committee report; scored CSV downloads"])]
    for i, (title, accent, items) in enumerate(gov):
        left = M + i * (cw + Inches(0.2))
        card(s, left, Inches(1.75), cw, Inches(2.6), title, "", accent=accent, title_size=14)
        bullets(s, left + Inches(0.2), Inches(2.2), cw - Inches(0.35), Inches(2.1), items, size=10.5)
    shots = [p for p in (MEDIA / "shap_top.png", MEDIA / "warnings.png", MEDIA / "drift.png") if p.exists()]
    if shots:
        sw = (W - 2 * M - Inches(0.2) * (len(shots) - 1)) / len(shots)
        for i, p in enumerate(shots):
            d.picture(s, p, M + i * (sw + Inches(0.2)), Inches(4.55), width=sw, height=Inches(2.15))

    # 12 · Scoring -------------------------------------------------------------
    s = d.slide("The trained model scores the next application from a form or a batch file, with columns harmonised automatically",
                "Scoring new applications", section="Product")
    bullets(s, M, Inches(1.75), Inches(5.4), Inches(4.9),
            [("Single application  ", "form generated from the model's input schema; gauge, band, probability, base vs AI-adjusted score, reason codes"),
             ("Batch file  ", "drop a CSV; exact → normalised → LLM column mapping shown with a method badge per column; values normalised (1,20,000 · 2.5 lakh · yes/no)"),
             ("Drift  ", "score PSI and top drifting inputs versus the training population on every batch"),
             ("Output  ", "scored CSV with nu_score, nu_band, probability, base_score and reasons 1–4"),
             ("Where the model lives  ", "a persisted bundle per training job (preprocessing, champion, blend, calibrator, bands, SHAP background, drift profiles); the training data itself is deleted")],
            size=11)
    shot = MEDIA / ("scoring_result.png" if (MEDIA / "scoring_result.png").exists() else "scoring.png")
    d.picture(s, shot, M + Inches(5.7), Inches(1.7), width=W - 2 * M - Inches(5.7), height=Inches(2.6))
    d.picture(s, MEDIA / "scoring_form.png", M + Inches(5.7), Inches(4.45), width=W - 2 * M - Inches(5.7), height=Inches(2.3))

    # 13 · Privacy & commercials ------------------------------------------------
    s = d.slide("Privacy by construction: nothing is stored, no PII is needed, and every model runs on Northern Arc hardware",
                "Privacy, security and commercials", section="Governance")
    cw = (W - 2 * M - Inches(0.2)) / 2
    card(s, M, Inches(1.75), cw, Inches(3.35), "Privacy and security", "", accent=GOOD, title_size=15)
    bullets(s, M + Inches(0.2), Inches(2.25), cw - Inches(0.35), Inches(2.8),
            ["Uploaded files are deleted the moment the training job succeeds; abandoned sessions expire in 24 h",
             "No PII required; PII-looking columns are detected by name and value pattern and dropped",
             "Only aggregate metrics and the model bundle remain; WOE tables anonymise category labels",
             "Local language models only (Ollama on-prem); they see column names and summary statistics, never rows",
             "Internal network only at the reverse proxy; hardened systemd service (memory and CPU fences, no new privileges, private tmp)",
             "GPUs and caches on the 3 TB data volume; nothing leaves the host"], size=11)
    card(s, M + cw + Inches(0.2), Inches(1.75), cw, Inches(3.35), "Commercials and positioning", "", accent=CYAN, title_size=15)
    bullets(s, M + cw + Inches(0.4), Inches(2.25), cw - Inches(0.35), Inches(2.8),
            ["Free during the introductory period",
             "Thereafter ₹10 per case scored — compute on Northern Arc servers included",
             "Model building, validation, reports and Nu Score bands included",
             "Positioned as 'democratising ML modelling': loan default, collections priority, churn, sales forecasting, inventory demand, fraud flags",
             "Internal adoption first; partner NBFCs and originators next"], size=11)
    text(s, M, Inches(5.3), Inches(6), Inches(0.35), "Data lifecycle", size=13, bold=True, color=NAVY)
    _chevrons(s, Inches(5.7), [("Upload", "chunked, resumable, internal network"), ("Convert", "Parquet on the data volume"),
                               ("Train", "GPUs; local LLM sees metadata only"), ("Delete", "session data removed on success"),
                               ("Keep", "aggregate metrics + model bundle")])

    # 14 · Architecture --------------------------------------------------------
    s = d.slide("A single hardened service on existing infrastructure: FastAPI, DuckDB, PyTorch and local LLMs behind nginx",
                "Architecture and operations", section="Technology")
    boxes = [("Browser", "Single-page app · Plotly · no external CDNs", M, Inches(1.9), Inches(2.6)),
             ("nginx", "/nuscore_claudecode/ · internal CIDRs · 20 GB uploads", M + Inches(2.9), Inches(1.9), Inches(2.6)),
             ("FastAPI service", "sessions · jobs · models · reports · uvicorn :8096", M + Inches(5.8), Inches(1.9), Inches(3.0)),
             ("Training workers", "process pool · 16 threads/job · GPUs 2–3", M + Inches(9.1), Inches(1.9), Inches(3.0)),
             ("DuckDB / Polars", "streaming CSV→Parquet · profiling · joins", M + Inches(2.9), Inches(3.9), Inches(2.6)),
             ("Engine", "scikit-learn · XGBoost · LightGBM · CatBoost · Optuna · PyTorch · SHAP", M + Inches(5.8), Inches(3.9), Inches(3.0)),
             ("Ollama (local LLM)", "qwen2.5 · GPUs 0–1 · roles, mapping, narrative", M + Inches(9.1), Inches(3.9), Inches(3.0)),
             ("Disk (/mnt)", "sessions (deleted after training) · job results · model bundles · logs", M, Inches(3.9), Inches(2.6))]
    for title, body, left, top, wdt in boxes:
        rect(s, left, top, wdt, Inches(1.5), WASH if title in ("Engine", "Ollama (local LLM)", "Training workers") else GROUND, line=LINE)
        rect(s, left, top, wdt, Inches(0.07), CYAN)
        text(s, left + Inches(0.15), top + Inches(0.15), wdt - Inches(0.3), Inches(0.4), title, size=13, bold=True, color=NAVY)
        text(s, left + Inches(0.15), top + Inches(0.6), wdt - Inches(0.3), Inches(0.9), body, size=10.5, color=INK2)
    bullets(s, M, Inches(5.7), W - 2 * M, Inches(1.2),
            [("Host  ", "64 cores · 432 GB RAM · 4 × Tesla T4 · 3 TB data volume; two concurrent training jobs"),
             ("Operations  ", "systemd unit nuscore-otg · smoke and end-to-end tests in deploy/ · engine invariants in tests/ · repository github.com/harshitrathi14/ml-on-the-go")], size=10.5)

    # 15 · Roadmap -------------------------------------------------------------
    s = d.slide("Delivered in one sprint; the next steps deepen the science and widen the use cases", "Roadmap", section="Next")
    done = ["Chunked multi-file uploads, bureau JSON, identifier matching", "PII / leakage screening, IV & WOE, clean-up log",
            "Nine-family tuned zoo incl. FT-Transformer; stacked blend", "Five-cohort validation, selection-safe champion, warnings",
            "Nu Score bands, SHAP, reason codes, PSI drift, bootstrap CI", "Scoring form & batch with column harmonisation", "Reports, tests, hardened deployment"]
    nxt = ["Regression and multi-class outcomes (sales, demand, LGD)", "Aggregation windows for monthly bank / repayment rows",
           "Grouped validation for repeated borrowers; reject inference", "Monotonic constraints and scorecard export (points per bin)",
           "Monitoring dashboard: PSI over time, band migration, back-testing", "Sequence models on transaction histories; text embeddings for narrations",
           "Multi-user workspaces and audit trail when opened to partners"]
    cw = (W - 2 * M - Inches(0.3)) / 2
    card(s, M, Inches(1.75), cw, Inches(3.35), "Live today", "", accent=GOOD, title_size=15)
    bullets(s, M + Inches(0.2), Inches(2.25), cw - Inches(0.35), Inches(2.8), done, size=11)
    card(s, M + cw + Inches(0.3), Inches(1.75), cw, Inches(3.35), "Next", "", accent=CYAN, title_size=15)
    bullets(s, M + cw + Inches(0.5), Inches(2.25), cw - Inches(0.35), Inches(2.8), nxt, size=11)
    text(s, M, Inches(5.3), Inches(6), Inches(0.35), "Phases shipped", size=13, bold=True, color=NAVY)
    _chevrons(s, Inches(5.7), [("0 · Go live", "brand, jobs, nginx, systemd"), ("1 · Big data", "chunked uploads, joins, screening"),
                               ("2 · Engine", "tuned zoo, blend, champion"), ("3 · Nu Score", "calibration, bands, scoring"),
                               ("4 · Governance", "cohorts, drift, diagnostics, tests")])

    # 16 · Close ---------------------------------------------------------------
    s = d.slide(dark=True)
    rect(s, 0, 0, W, H, NAVY)
    d.logos(s, Inches(0.55))
    text(s, M, Inches(2.3), Inches(11), Inches(1.2), "Build ML models with your data.", size=40, bold=True, color=WHITE)
    text(s, M, Inches(3.4), Inches(10), Inches(1.2),
         "Upload a loan book today and score the next application in minutes — validated, explainable, on-prem.",
         size=18, color=RGBColor(0xC7, 0xD2, 0xE0))
    text(s, M, Inches(4.9), Inches(10), Inches(0.5), "http://10.91.16.86/nuscore_claudecode/", size=16, color=CYAN, bold=True)
    text(s, M, Inches(5.4), Inches(10), Inches(0.5), "Usage guide: docs/NU_SCORE_USAGE_GUIDE.md  ·  Free during introduction, then ₹10 per case",
         size=12, color=RGBColor(0xC7, 0xD2, 0xE0))

    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)
    print(f"wrote {out} ({d.n} slides)")


def latest_result() -> dict:
    root = Path(os.environ.get("DATA_ROOT", "/mnt/harshitrathi/mlpipeline-data")) / "jobs"
    candidates = sorted(root.glob("*/result.json"), key=lambda p: p.stat().st_mtime, reverse=True)
    for path in candidates:
        data = json.loads(path.read_text())
        if data.get("nuscore"):
            return data
    return {}


if __name__ == "__main__":
    result = json.loads(Path(sys.argv[1]).read_text()) if len(sys.argv) > 1 else latest_result()
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "docs" / "Nu_Score_On_the_Go.pptx"
    build(result, out)
